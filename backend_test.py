socketserver_reuse=True
import json, os, subprocess, sys, time, threading, socketserver, urllib.request, urllib.error, zipfile, io, tempfile, http.cookiejar, sqlite3

ROOT=os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "server")
DATA=tempfile.mkdtemp(prefix="apis_test_")
PORT=18501
import datetime
T=datetime.date.today().isoformat(); Y=(datetime.date.today()-datetime.timedelta(days=1)).isoformat(); W=(datetime.date.today()-datetime.timedelta(days=6)).isoformat()
proc=subprocess.Popen([sys.executable, ROOT+"/apis_server.py","--no-browser","--data",DATA,"--port",str(PORT)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT)
BASE="http://127.0.0.1:%d"%PORT
cj=http.cookiejar.CookieJar()
op=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
def call(method,path,body=None,raw=False,headers=None):
    h={"X-APIS":"1"}; data=None
    if body is not None and not isinstance(body,(bytes,bytearray)):
        data=json.dumps(body).encode(); h["Content-Type"]="application/json"
    elif body is not None: data=body
    h.update(headers or {})
    req=urllib.request.Request(BASE+path,data=data,method=method,headers=h)
    try:
        r=op.open(req); b=r.read()
        return r.status,(b if raw else json.loads(b.decode() or "null"))
    except urllib.error.HTTPError as e:
        b=e.read()
        try: return e.code,json.loads(b.decode())
        except Exception: return e.code,b
def ok(c,msg):
    print(("PASS " if c else "FAIL ")+msg)
    if not c: sys.exit(1)
for i in range(50):
    try: urllib.request.urlopen(BASE+"/api/ping",timeout=1); break
    except Exception: time.sleep(0.2)
try:
    s,r=call("GET","/api/ping"); ok(s==200 and r["app"]=="APIS","ping")
    ok(call("GET","/api/dashboard")[0]==401,"api requires login")
    ok(call("GET","/api/auth/status")[1]["setup"] is False,"fresh install has no pin")
    ok(call("POST","/api/auth/setup",{"pin":"12"})[0]==400,"short pin rejected")
    ok(call("POST","/api/auth/setup",{"pin":"2468"})[0]==200,"pin set")
    s,r=call("GET","/api/bootstrap"); ok(r["profile"]["name"]!="","starter profile exists")
    ok(call("PUT","/api/profile",{"name":"Test Person"})[1]["name"]=="Test Person","profile edited")
    call("POST","/api/experience",{"org":"Org A","role":"Old","start":"2020-01","end":"2021-01"}); call("POST","/api/experience",{"org":"Org B","role":"New","start":"2024-01"})
    ex=call("GET","/api/experience")[1]; ok(ex[0]["org"]=="Org B" and ex[-1]["org"]=="Org A","experience ongoing first, oldest last")
    call("POST","/api/auth/logout"); ok(call("GET","/api/dashboard")[0]==401,"logout works")
    ok(call("POST","/api/auth/login",{"pin":"0000"})[0]==401,"wrong pin")
    ok(call("POST","/api/auth/login",{"pin":"2468"})[0]==200,"login")
    s,r=call("POST","/api/logs",{"date":T,"category":"Task","title":"","impact":"High","status":"Completed"}); ok(s==400 and "required" in r["error"],"log validation")
    s,l1=call("POST","/api/logs",{"date":T,"category":"Field visit","title":"Retailer visit Mirpur","impact":"High","status":"Completed","hours":1.5,"result":"2 new orders","location":"Mirpur-10"}); ok(s==201,"create log")
    s,l2=call("POST","/api/logs",{"date":T,"category":"Admin","title":"Update sheet","impact":"Low","status":"Pending"}); ok(s==201,"create log 2")
    ok(len(call("GET","/api/appraisal")[1])==1,"appraisal only for completed medium/high")
    call("PUT","/api/logs/"+l2["id"],{"status":"Completed","impact":"Medium"}); ok(len(call("GET","/api/appraisal")[1])==2,"appraisal added on completion")
    call("PUT","/api/logs/"+l2["id"],{"title":"Renamed"}); ok(len(call("GET","/api/appraisal")[1])==2,"no duplicate appraisal rows")
    ok(call("GET","/api/appraisal/verify")[1]["ok"],"chain verifies")
    c=sqlite3.connect(DATA+"/apis.db")
    try: c.execute("UPDATE appraisal SET title='x'"); ok(False,"trigger should block")
    except sqlite3.DatabaseError as e: ok("cannot be edited" in str(e),"db blocks edits to career log")
    c.execute("DROP TRIGGER appraisal_no_update"); c.execute("UPDATE appraisal SET title='tampered' WHERE seq=1"); c.commit(); c.close()
    v=call("GET","/api/appraisal/verify")[1]; ok(v["ok"] is False and v["broken_at"]==1,"tampering detected")
    for d,route,t,a_,so in [(T,"Route 07",120000,94000,1),(T,"Route 08",80000,85000,0),(Y,"Route 07",100000,70000,0)]:
        ok(call("POST","/api/territory",{"date":d,"territory":"Dhaka North","route":route,"distributor":"ABC Traders","target":t,"actual":a_,"stock_out":bool(so),"stock_out_items":"Product A 500g" if so else ""})[0]==201,"territory "+route)
    s,sm=call("GET","/api/territory/summary?from=%s&to=%s"%(W,T)); ok(len(sm["routes"])==2 and sm["target"]==300000,"territory summary")
    ok(sm["routes"][0]["pct"]==round(164000/220000*100,1),"route pct maths")
    d=call("GET","/api/dashboard")[1]; ok(d["tasks"]["total"]==2 and any(a["kind"]=="stock" for a in d["alerts"]),"dashboard alerts")
    s,ld=call("POST","/api/leads",{"company":"Sample Traders Ltd","contact":"Mr Rahman","product":"SME loan","amount":2500000,"next_action":"Collect financials","next_action_date":"2026-10-10"}); ok(s==201 and ld["stage"]=="Prospect","create lead")
    docs=call("GET","/api/lead_docs?lead_id="+ld["id"])[1]; ok(len(docs)==6,"6 default docs")
    ok(call("POST","/api/leads/%s/move"%ld["id"],{"stage":"Doc Verification"})[0]==200,"move to doc verification")
    s,r=call("POST","/api/leads/%s/move"%ld["id"],{"stage":"Underwriting"}); ok(s==409 and "Trade License" in r["missing"],"gate blocks underwriting")
    for d_ in docs:
        if d_["required"]: call("PUT","/api/lead_docs/"+d_["id"],{"received":True,"verified":True})
    ok(call("POST","/api/leads/%s/move"%ld["id"],{"stage":"Underwriting"})[0]==200,"gate opens when verified")
    ok(any(a["kind"]=="lead" for a in call("GET","/api/dashboard")[1]["alerts"]),"overdue follow-up alert")
    rp=call("GET","/api/report/preview?date="+T)[1]; ok("Retailer visit Mirpur" in rp["text"] and "STOCK-OUT" in rp["text"],"report preview text")
    s,x=call("GET","/api/export.xlsx?scope=all",raw=True); open(DATA+"/all.xlsx","wb").write(x)
    import openpyxl, warnings; warnings.simplefilter("ignore")
    wb=openpyxl.load_workbook(DATA+"/all.xlsx"); ok("Daily_Logs" in wb.sheetnames and wb["Daily_Logs"].max_row==3,"xlsx opens: %s"%wb.sheetnames)
    s,x=call("GET","/api/export.xlsx?scope=today&date="+T,raw=True); ok(openpyxl.load_workbook(io.BytesIO(x))["Territory_Metrics"].max_row==3,"today xlsx")
    from PIL import Image
    b=io.BytesIO(); Image.new("RGB",(40,30),(10,80,200)).save(b,"JPEG")
    s,u=call("POST","/api/upload",b.getvalue(),headers={"Content-Type":"image/jpeg"}); ok(s==201 and u["file"].endswith(".jpg"),"upload jpg")
    ok(call("POST","/api/upload",b"not an image",headers={"Content-Type":"image/jpeg"})[0]==400,"upload rejects non-images")
    s,p=call("POST","/api/posts",{"caption":"Visited the Sales Carnival","photos":[u["file"]],"visibility":"public"}); ok(s==201 and p["photos"]==[u["file"]],"post with photo")
    ok(call("POST","/api/posts",{"caption":"","photos":[]})[0]==400,"empty post rejected")
    s,im=call("GET","/files/"+u["file"],raw=True); ok(s==200 and im[:2]==b"\xff\xd8","serve file")
    ok(call("POST","/api/certs",{"title":"Test cert","issuer":"X","issued":"2026-10","image":u["file"]})[0]==201,"cert with image")
    s,z=call("GET","/api/publish.zip",raw=True); zf=zipfile.ZipFile(io.BytesIO(z)); ok("index.html" in zf.namelist() and any(n.startswith("img/") for n in zf.namelist()),"public site zip")
    html=zf.read("index.html").decode(); ok("Visited the Sales Carnival" in html and "<script" not in html,"public html has public post, no script")
    ok(b"Test Person" in call("GET","/cv",raw=True)[1],"cv page")
    ok(call("DELETE","/api/posts/"+p["id"])[0]==200 and os.path.exists(DATA+"/uploads/"+u["file"]),"delete post, photo kept while still used by cert")
    ok(call("POST","/api/backup/run")[1]["local"],"local backup")
    drive=tempfile.mkdtemp(prefix="fakedrive_")
    call("PUT","/api/settings",{"drive_dir":drive}); b=call("POST","/api/backup/run")[1]
    ok(b["drive"] and os.path.exists(drive+"/APIS Backup/apis_data.xlsx") and not os.path.exists(drive+"/APIS Backup/secrets.json"),"drive backup copies db+xlsx, no secrets")
    call("PUT","/api/settings",{"drive_dir":"/nope/none"}); b=call("POST","/api/backup/run")[1]; ok(not b["drive"] and "not found" in b["drive_error"],"drive error message")
    ok(call("PUT","/api/settings",{"eod_time":"25:99"})[0]==400,"bad time rejected")
    ok(call("PUT","/api/settings",{"report_to":"nope"})[0]==400,"bad email rejected")
    msgs=[]
    class SMTP(socketserver.StreamRequestHandler):
        def handle(self):
            w=lambda s: self.wfile.write((s+"\r\n").encode()); w("220 fake ESMTP"); data=False; buf=[]
            while True:
                line=self.rfile.readline()
                if not line: break
                t=line.decode(errors="replace")
                if data:
                    if t in (".\r\n",".\n"): msgs.append("".join(buf)); data=False; buf=[]; w("250 queued")
                    else: buf.append(t)
                    continue
                c=t.strip().upper()
                if c.startswith(("EHLO","HELO")): w("250-fake"); w("250 AUTH PLAIN LOGIN")
                elif c.startswith("AUTH"): w("235 ok")
                elif c.startswith(("MAIL","RCPT")): w("250 ok")
                elif c.startswith("DATA"): w("354 go"); data=True
                elif c.startswith("QUIT"): w("221 bye"); break
                else: w("250 ok")
    socketserver.ThreadingTCPServer.allow_reuse_address=True; srv=socketserver.ThreadingTCPServer(("127.0.0.1",12525),SMTP); srv.daemon_threads=True
    threading.Thread(target=srv.serve_forever,daemon=True).start()
    s,r=call("POST","/api/report/send",{"date":T}); ok(s==400 and "Connect Gmail" in r["error"],"send blocked until gmail set")
    s,r=call("PUT","/api/gmail",{"address":"me@gmail.com","app_password":"abcd efgh ijkl mnop","report_to":"boss@example.com, me@gmail.com"}); ok(s==200 and r["gmail_password_set"] and "abcdefgh" not in json.dumps(r),"gmail saved, password not echoed")
    ok(json.load(open(DATA+"/secrets.json"))["gmail_app_password"]=="abcdefghijklmnop","password stored without spaces")
    call("PUT","/api/settings",{"smtp_host":"127.0.0.1","smtp_port":12525,"smtp_ssl":False})
    s,r=call("POST","/api/report/send",{"date":T}); ok(s==200 and r["ok"],"report sent via smtp")
    ok(len(msgs)==1 and "boss@example.com" in msgs[0] and 'filename="APIS_EOD_%s.xlsx"'%T in msgs[0] and "Retailer visit Mirpur" in msgs[0],"email has recipients, attachment and body")
    ok(call("POST","/api/mail/test")[0]==200,"test email")
    ok(len(call("GET","/api/sends")[1])==1,"send history")
    call("PUT","/api/settings",{"smtp_port":12599}); s,r=call("POST","/api/report/send",{"date":T}); ok(s==400 and "could not reach" in r["error"].lower(),"unreachable server message")
    ok(call("POST","/api/auth/change",{"old":"2468","new":"13579"})[0]==200,"pin changed")
    req=urllib.request.Request(BASE+"/api/logs",data=b"{}",method="POST",headers={"Content-Type":"application/json"})
    try: op.open(req); ok(False,"csrf")
    except urllib.error.HTTPError as e: ok(e.code==403,"POST without X-APIS header blocked")
    ok(call("GET","/..%2f..%2fdata%2fapis.db",raw=True)[0]==404,"path traversal blocked")
    ok(call("GET","/server/store.py",raw=True)[0]==404,"python source not served")
    print("\nALL BACKEND TESTS PASSED")
finally:
    proc.terminate()
    log=os.path.join(DATA,"apis.log")
    if os.path.exists(log):
        t=open(log).read()
        if "Traceback" in t: print("SERVER LOG TRACEBACKS:\n"+t[-3000:])
